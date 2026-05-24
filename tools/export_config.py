#!/usr/bin/env python3
"""Excel -> JSON export for Godot runtime.

Normal usage (after editing Excel):
    python tools/export_config.py

First-time / reset Excel templates from defaults:
    python tools/export_config.py --init-excel
"""
from __future__ import annotations

import json
import sys
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill

ROOT = Path(__file__).resolve().parents[1]
EXCEL_DIR = ROOT / "config" / "excel"
JSON_DIR = ROOT / "config" / "json"

HEADER_FILL = PatternFill("solid", fgColor="4472C4")
HEADER_FONT = Font(color="FFFFFF", bold=True)

EXCEL_SHEETS = [
    "chapters",
    "stages",
    "monsters",
    "player",
    "upgrades",
    "upgrade_fx",
    "game_tuning",
    "asset_mapping",
]

CHAPTER_HEADERS = [
    "chapter_id", "chapter_name", "stages_per_chapter", "description",
]
CHAPTER_ROWS = [
    [1, "第一章", 4, "入门章节"],
    [2, "第二章", 4, "Boss 与高强度混合战"],
]

STAGE_HEADERS = [
    "stage_index", "chapter_id", "stage_in_chapter", "display_name",
    "normal", "elite", "shield", "berserker", "splitter", "archer", "fire_mage", "boss_id",
]
STAGE_ROWS = [
    [0, 1, 1, "第1关", 62, 0, 0, 0, 0, 0, 0, ""],
    [1, 1, 2, "第2关", 60, 12, 0, 0, 0, 10, 0, ""],
    [2, 1, 3, "第3关", 64, 20, 8, 0, 0, 12, 6, ""],
    [3, 1, 4, "第4关", 64, 20, 12, 8, 0, 14, 7, ""],
    [4, 2, 1, "第5关", 0, 0, 0, 0, 0, 0, 0, "centipede"],
    [5, 2, 2, "第6关", 68, 24, 16, 12, 0, 18, 9, ""],
    [6, 2, 3, "第7关", 72, 28, 20, 12, 0, 20, 10, ""],
    [7, 2, 4, "第8关", 76, 28, 24, 16, 0, 22, 11, ""],
]

MONSTER_HEADERS = [
    "kind_id", "name_cn", "hp", "def", "attack", "attack_interval", "size", "speed",
    "color_hex", "grade", "can_move", "attack_range", "arrow_speed", "ranged",
    "ki_drain_on_hit", "max_split_tier", "split_count", "exp_reward",
    "character_folder", "sprite_prefix", "projectile_folder",
]
MONSTER_ROWS = [
    ["NORMAL", "普通怪", 80, 4, 10, 1.15, 13, 19, "#526078", "B", 1, 0, 0, 0, 0, 0, 0, 2, "Skeleton", "Skeleton", ""],
    ["ELITE", "高级怪", 110, 6, 14, 1.05, 14, 21, "#8a6aa8", "B+", 1, 0, 0, 0, 0, 0, 0, 4, "Elite Orc", "Elite Orc", ""],
    ["SHIELD", "盾牌怪", 120, 5, 12, 1.25, 15, 16, "#5f7a88", "B+", 1, 0, 0, 0, 0, 0, 0, 3, "Knight", "Knight", ""],
    ["BERSERKER", "狂战士", 90, 2, 18, 0.85, 14, 26, "#b85a4a", "C/B", 1, 0, 0, 0, 10, 0, 0, 4, "Armored Axeman", "Armored Axeman", ""],
    ["SPLITTER", "分裂怪", 70, 2, 8, 1.2, 18, 14, "#7b9f5a", "A+/C/C", 1, 0, 0, 0, 0, 3, 2, 3, "Slime", "Slime", ""],
    ["ARCHER", "射箭怪", 75, 3, 12, 1.35, 12, 17, "#6a8a5a", "B", 1, 215, 85, 1, 0, 0, 0, 3, "Archer", "Archer", "Arrow(projectile)"],
    ["FIRE_MAGE", "火焰法师", 70, 2, 16, 2.4, 12, 15, "#501818", "B", 1, 250, 0, 1, 0, 0, 0, 3, "Wizard", "Wizard", "Magic(projectile)"],
]

PLAYER_HEADERS = ["key", "value", "description"]
PLAYER_ROWS = [
    ["base_attack", 95, "基础攻击力"],
    ["base_hp", 100, "基础生命"],
    ["base_ki", 234, "基础气力上限"],
    ["base_crit_rate", 0.08, "暴击率"],
    ["base_crit_damage", 1.6, "暴击伤害倍率"],
    ["attack_speed", 2300, "路径冲刺速度 px/s"],
    ["hitbox_radius", 12, "碰撞半径"],
    ["trigger_radius_ratio", 0.06, "触发圈相对屏宽比例"],
    ["trigger_radius_min", 30, "触发圈最小半径"],
    ["ki_per_pixel", 0.18, "划线每像素消耗气力"],
    ["ki_regen_rate", 135, "气力恢复速度/s"],
    ["combo_damage_bonus", 0.01, "每连击伤害加成"],
    ["invincible_time", 0.45, "受击无敌时间"],
    ["auto_dart_interval", 0.5, "自动飞镖间隔"],
    ["auto_dart_damage_mult", 0.20, "自动飞镖伤害倍率"],
    ["auto_dart_speed", 420, "自动飞镖速度"],
    ["auto_dart_life", 0.9, "自动飞镖存活时间"],
    ["character_folder", "Swordsman", "主角素材文件夹"],
    ["sprite_prefix", "Swordsman", "主角精灵前缀"],
    ["sprite_scale", 2, "显示缩放(100px素材，2=200px显示)"],
]

UPGRADE_HEADERS = [
    "id", "name_cn", "rarity", "icon", "desc_cn", "is_pet", "max_level",
    "apply_type", "apply_value", "trigger_condition", "effect_pack", "effect_name",
]
UPGRADE_ROWS = [
    ["luck", "运气", "white", "🧘", "增加20%气力", 0, 9, "ki_mult", 1.20, "", "", ""],
    ["multi_dart", "多重飞镖", "white", "🎯", "普攻飞镖数量+1", 0, 9, "dart_count", 1, "", "effects/3.Gothicvania Magic Pack N2", "dart"],
    ["giant_dart", "巨大飞镖", "blue", "⭕", "普攻20%概率发射巨大飞镖", 0, 9, "giant_dart_chance", 0.20, "", "effects/3.Gothicvania Magic Pack N2", "giant_dart"],
    ["ice_dart", "寒冰飞镖", "purple", "❄️", "普攻命中5%释放寒冰飞镖", 0, 9, "ice_dart_chance", 0.05, "", "effects/2.Gothicvania Magic Pack N3", "ice"],
    ["spirit_bomb", "元气弹", "orange", "💫", "飞镖变元气弹", 0, 9, "spirit_bomb", 1, "", "effects/10.GothicVania Magic Pack 7", "spirit_bomb"],
    ["shuriken", "手里剑", "blue", "🎯", "每次连击释放2枚手里剑", 0, 9, "shuriken_count", 2, "combo_hit", "effects/3.Gothicvania Magic Pack N2", "shuriken"],
    ["multi_combo", "多重连击", "orange", "🔥", "连击次数×1.2倍", 0, 9, "combo_mult", 1.20, "", "", ""],
    ["healing_combo", "愈合连击", "orange", "🌿", "连击每+15释放藤蔓，回复5%生命", 0, 9, "heal_pct", 0.05, "combo_milestone_15", "effects/4.Gothicvania Magic Pack N4", "vine"],
    ["holy_shield", "圣盾", "blue", "🛡️", "每5秒获得1层护盾", 0, 9, "holy_shield_interval", 5.0, "passive", "effects/5.Gothicvania Magic Pack N5", "shield"],
    ["vampire_bat", "吸血蝙蝠", "purple", "🦇", "每击杀10敌人，蝙蝠群攻并回复2%生命", 0, 9, "heal_pct", 0.02, "kill_10", "effects/6.Gothicvania Magic Pack N6", "bat"],
    ["meat_shield", "变肉", "white", "🛡️", "最大生命+10%，体积+15%", 0, 9, "hp_mult", 1.10, "", "", ""],
    ["super_heal", "超级治疗", "orange", "✨", "所有回血效果+50%", 0, 9, "heal_bonus", 0.50, "", "effects/7.Gothicvania Magic Pack N7", "heal"],
    ["lightning_chain", "闪电链", "blue", "⚡", "连击每+5释放闪电链", 0, 9, "chain", 1, "combo_milestone_5", "effects/8.Gothicvania Magic Pack N8", "lightning"],
    ["shadow_clone", "影分身", "blue", "👤", "连击>10召唤影分身", 0, 9, "clone", 1, "combo_gt_10", "effects/9.Gothicvania Magic Pack N9", "clone"],
    ["water_tornado", "水龙卷术", "purple", "🌊", "暴击+5%，连击每+3释放水龙卷", 0, 9, "crit_rate", 0.05, "combo_milestone_3", "effects/13.Gothicvania Magic Pack 8/Magic Pack 8 files/sprites", "tornado"],
    ["black_hole", "黑洞", "purple", "🕳️", "连击8时生成黑洞", 0, 9, "black_hole", 1, "combo_eq_8", "effects/11.Warped Explosions Pack 6", "black_hole"],
    ["blade_whirl", "刀阵旋风", "purple", "🌀", "连击每+5释放刀阵旋风", 0, 9, "whirl", 1, "combo_milestone_5", "effects/12.Gothicvania Magic Pack 8", "whirl"],
    ["great_fireball", "豪火球术", "orange", "🔥", "连击每+10释放豪火球", 0, 9, "fireball", 1, "combo_milestone_10", "effects/1.Gothicvania Magic Pack N1", "fireball"],
    ["heavenly_thunder", "天雷", "orange", "⚡", "每秒落雷范围攻击", 0, 9, "thunder", 1, "passive_1s", "effects/8.Gothicvania Magic Pack N8", "thunder"],
    ["wild_wolf", "野狼", "white", "🐺", "宠物：召唤一只野狼", 1, 9, "pet_wolf", 1, "", "Characters/Characters(100x100)/Werewolf", "Werewolf"],
    ["wild_bull", "野牛", "blue", "🐂", "宠物：召唤一只野牛", 1, 9, "pet_bull", 1, "", "Characters/Characters(100x100)/Werebear", "Werebear"],
    ["divine_god", "天神", "purple", "✨", "宠物：召唤天神", 1, 9, "pet_god", 1, "", "Characters/Characters(100x100)/Knight Templar", "Knight Templar"],
    ["nurturing_heart", "养育之心", "orange", "💗", "宠物数量翻倍", 0, 9, "pet_mult", 2, "", "", ""],
]

UPGRADE_FX_HEADERS = [
    "rarity", "name_cn", "chance", "color_hex", "overlay", "edge_glow", "card_glow",
    "shimmer", "pulse", "spark_count", "rays",
]
UPGRADE_FX_ROWS = [
    ["white", "普通", 0.30, "#d8d8d8", 0.76, 0, 0.12, 0, 0, 0, 0],
    ["blue", "稀有", 0.30, "#58a8ff", 0.80, 0.28, 0.32, 1, 0, 6, 0],
    ["purple", "史诗", 0.30, "#b070ff", 0.84, 0.48, 0.50, 1, 1, 12, 0],
    ["orange", "传奇", 0.10, "#ff9830", 0.88, 0.78, 0.85, 1, 1, 22, 1],
]

TUNING_HEADERS = ["key", "value", "description"]
TUNING_ROWS = [
    ["logical_width", 390, "逻辑宽度"],
    ["logical_height", 700, "逻辑高度"],
    ["unit_scale", 1.3, "单位缩放"],
    ["camera_zoom", 1.0, "战斗镜头缩放(1=正常，勿随意改大)"],
    ["monster_sprite_scale", 2, "怪物显示缩放"],
    ["bullet_time_scale", 0.14, "子弹时间缩放"],
    ["bullet_time_dim_alpha", 0.42, "子弹时间暗化"],
    ["stages_per_chapter", 4, "每章关卡数"],
    ["stage_hp_growth", 1.3, "关卡HP缩放幂次"],
    ["stage_def_growth", 1.3, "关卡DEF缩放幂次"],
    ["stage_monster_scale", 1.3, "关卡怪物数量缩放"],
    ["stage_count_mul", 0.666667, "怪物数量乘数"],
    ["shield_count_mul", 0.333333, "盾牌怪数量乘数"],
    ["exp_base_to_level", 100, "升级所需基础经验"],
    ["exp_growth", 1.22, "升级经验增长"],
    ["terrain_folder", "Terrain", "地形素材目录"],
    ["ui_folder", "ui", "UI素材目录"],
    ["effects_folder", "effects", "特效素材目录"],
    ["fire_pillar_radius", 60, "火柱半径"],
    ["fire_pillar_warning_time", 1.1, "火柱预警时间"],
    ["fire_pillar_active_time", 0.5, "火柱持续时间"],
    ["fire_pillar_fade_time", 0.3, "火柱淡出时间"],
    ["stage_intro_slide_in", 0.38, "关卡 intro 滑入时长"],
    ["stage_intro_hold", 0.85, "关卡 intro 停留时长"],
    ["stage_intro_slide_out", 0.38, "关卡 intro 滑出时长"],
    ["stage_intro_sakura_extra", 0.8, "intro 樱花额外停留"],
    ["stage_clear_flash_duration", 1.1, "过关闪白时长"],
    ["stage_fail_label_duration", 1.2, "失败字幕时长"],
    ["stage_fail_overlay_fade", 0.65, "失败遮罩淡入时长"],
    ["fail_death_windup", 0.3, "失败动画蓄力"],
    ["fail_death_spear_stagger", 0.065, "失败投矛错峰"],
    ["fail_death_spear_speed", 640, "失败投矛速度"],
    ["fail_death_impact_pause", 0.24, "失败插矛停顿"],
    ["fail_death_duration", 0.9, "失败倒地时长"],
    ["fail_death_playback_speed", 2.0, "失败动画加速"],
    ["combat_first_hit_delay", 0.04, "连击首击延迟"],
    ["combat_hit_interval", 0.012, "连击命中间隔"],
    ["combat_afterimage_life", 0.1, "连击残影寿命"],
    ["combat_death_stagger", 0.012, "死亡错峰间隔"],
    ["monster_spawn_anim", 0.42, "怪物出生动画时长"],
    ["sakura_petal_count", 42, "樱花花瓣数量"],
    ["sakura_default_duration", 5.5, "樱花默认时长"],
    ["sakura_fade_out_duration", 1.8, "樱花淡出时长"],
    ["grass_cluster_max", 22, "草簇最大数量"],
    ["terrain_prop_count", 8, "地形装饰数量"],
    ["target_fps", 60, "目标帧率(0=不限制)"],
]

ASSET_HEADERS = ["asset_key", "category", "folder_path", "file_pattern", "notes"]
ASSET_ROWS = [
    ["terrain_bg", "terrain", "Terrain", "*.png", "战斗背景/地形"],
    ["ui_hud", "ui", "ui", "*.png", "HUD界面元素"],
    ["effect_magic", "effect", "effects", "*.png", "魔法特效序列帧"],
    ["character_base", "character", "Characters/Characters(100x100)", "*/*/*.png", "角色动画精灵表"],
]

DEFAULT_DATASETS = {
    "chapters": (CHAPTER_HEADERS, CHAPTER_ROWS),
    "stages": (STAGE_HEADERS, STAGE_ROWS),
    "monsters": (MONSTER_HEADERS, MONSTER_ROWS),
    "player": (PLAYER_HEADERS, PLAYER_ROWS),
    "upgrades": (UPGRADE_HEADERS, UPGRADE_ROWS),
    "upgrade_fx": (UPGRADE_FX_HEADERS, UPGRADE_FX_ROWS),
    "game_tuning": (TUNING_HEADERS, TUNING_ROWS),
    "asset_mapping": (ASSET_HEADERS, ASSET_ROWS),
}


def write_sheet(ws, headers: list[str], rows: list[list]):
    ws.append(headers)
    for cell in ws[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
    for row in rows:
        ws.append(row)
    for col in ws.columns:
        max_len = max(len(str(c.value or "")) for c in col)
        ws.column_dimensions[col[0].column_letter].width = min(max_len + 2, 40)


def normalize_cell(value):
    if value is None:
        return None
    if isinstance(value, bool):
        return value
    if isinstance(value, (int, float)):
        return value
    text = str(value).strip()
    if text == "":
        return ""
    if text.lower() in ("true", "false"):
        return text.lower() == "true"
    try:
        if "." in text:
            return float(text)
        return int(text)
    except ValueError:
        return text


def rows_to_dicts(headers: list[str], rows: list[list]) -> list[dict]:
    out = []
    for row in rows:
        item = {}
        for i, h in enumerate(headers):
            if not h:
                continue
            val = normalize_cell(row[i] if i < len(row) else None)
            item[h] = "" if val is None else val
        if any(str(v).strip() != "" for v in item.values() if v is not None):
            out.append(item)
    return out


def save_json(name: str, data):
    JSON_DIR.mkdir(parents=True, exist_ok=True)
    path = JSON_DIR / f"{name}.json"
    with path.open("w", encoding="utf-8") as f:
        json.dump(data, f, ensure_ascii=False, indent=2)
    print(f"  -> {path}")


def read_excel(name: str) -> tuple[list[str], list[list]]:
    path = EXCEL_DIR / f"{name}.xlsx"
    if not path.exists():
        raise FileNotFoundError(f"找不到 Excel 文件: {path}")
    wb = load_workbook(path, data_only=True)
    ws = wb.active
    raw_rows = [list(row) for row in ws.iter_rows(values_only=True)]
    if not raw_rows:
        return [], []
    headers = [str(h).strip() if h is not None else "" for h in raw_rows[0]]
    rows = []
    for raw in raw_rows[1:]:
        if all(cell is None or str(cell).strip() == "" for cell in raw):
            continue
        rows.append(list(raw))
    return headers, rows


def build_workbooks():
    EXCEL_DIR.mkdir(parents=True, exist_ok=True)
    for name, (headers, rows) in DEFAULT_DATASETS.items():
        wb = Workbook()
        ws = wb.active
        ws.title = name
        write_sheet(ws, headers, rows)
        path = EXCEL_DIR / f"{name}.xlsx"
        wb.save(path)
        print(f"Excel: {path}")


def merge_tuning_excel() -> int:
    """Append missing tuning keys from TUNING_ROWS without overwriting existing Excel values."""
    path = EXCEL_DIR / "game_tuning.xlsx"
    if not path.exists():
        return 0
    _, default_rows = DEFAULT_DATASETS["game_tuning"]
    wb = load_workbook(path)
    ws = wb.active
    existing_keys: set[str] = set()
    for row in ws.iter_rows(min_row=2, values_only=True):
        if row and row[0] is not None and str(row[0]).strip():
            existing_keys.add(str(row[0]).strip())
    added = 0
    for row in default_rows:
        key = str(row[0]).strip()
        if key and key not in existing_keys:
            ws.append(row)
            added += 1
    if added:
        wb.save(path)
        print(f"  game_tuning: merged {added} missing key(s) into Excel")
    return added


def export_json_from_excel():
    JSON_DIR.mkdir(parents=True, exist_ok=True)
    merge_tuning_excel()
    print("Reading Excel files...")
    for name in EXCEL_SHEETS:
        headers, rows = read_excel(name)
        data = rows_to_dicts(headers, rows)
        save_json(name, data)
        print(f"  {name}: {len(data)} rows")

    boss = {
        "centipede": {
            "name": "千足虫",
            "warning_time": 3,
            "segment_hp": 320,
            "hp_scale": 3.5,
            "segment_def": 3,
            "segment_radius": 23,
            "crawl_speed": 240,
            "bullet_interval": 0.38,
            "bullets_per_shot": 20,
            "bullet_damage": 24,
            "bullet_speed": 130,
            "defeat_exp": 140,
        }
    }
    save_json("bosses", boss)

    buff_orbs = {
        "base_types": ["attack", "ki", "combo"],
        "radius": 13,
        "max_per_type": 4,
        "spawn_chance": {"attack": 0.55, "ki": 0.45, "combo": 0.42, "ice": 0.28},
        "extra_spawn_chance": 0.48,
        "min_ki_per_stage": 2,
    }
    save_json("buff_orbs", buff_orbs)


def main():
    if "--init-excel" in sys.argv:
        print("Resetting Excel templates from defaults...")
        build_workbooks()
        print("Excel templates created.")
    export_json_from_excel()
    print("Done. Restart Godot game (F5) to apply changes.")


if __name__ == "__main__":
    main()
